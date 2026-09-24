"""Convertit les référentiels xlsx de data/ en JSON pour le site statique.

Usage :

    pip install -r scripts/requirements.txt
    python scripts/build_data.py

Un seul rôle : lire chaque xlsx déclaré dans JEUX et écrire son JSON dans
docs/assets/data/<theme>/<fichier>.json. Seul openpyxl est nécessaire pour
lire le xlsx, et rien d'autre que la bibliothèque standard pour écrire le
JSON.

Remplacer un xlsx dans data/, relancer ce script, committer le xlsx et le
JSON regénéré : c'est tout ce qu'il faut pour mettre à jour un référentiel.

L'arrêté tarifaire (data/groupage/tarifs.xlsx) est relu tel que l'ATIH le
publie : le script y prend la feuille des GHS du secteur public, en vérifie
l'en-tête et chaque ligne, et s'arrête plutôt que de deviner — un intitulé
déplacé, un couple GHS-GHM en double, un montant illisible ou une racine
inconnue de data/groupage/racines.xlsx arrêtent la conversion.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from millesime import millesime

RACINE = Path(__file__).resolve().parent.parent
DOSSIER_DONNEES = RACINE / "data"
DOSSIER_SORTIE = RACINE / "docs" / "assets" / "data"


@dataclass(frozen=True)
class Jeu:
    theme: str  # sous-dossier commun à data/<theme>/ et docs/assets/data/<theme>/
    fichier: str  # nom du xlsx (ou csv) dans data/<theme>/
    feuille: str | None = None  # feuille du classeur à lire ; None : la feuille active


# La seule liste à tenir à jour : ajouter un référentiel, c'est ajouter une
# ligne ici (et la déclarer dans assets/js/registry.js côté site).
JEUX: tuple[Jeu, ...] = (
    Jeu("intox", "medicaments.xlsx"),
    Jeu("intox", "substances.xlsx"),
    Jeu("germes", "germes.xlsx"),
    Jeu("actes", "actes.xlsx"),
    Jeu("contextes", "contextes.xlsx"),
    Jeu("acronymes", "acronymes.xlsx"),
    Jeu("groupage", "diagnostics.xlsx"),
    Jeu("groupage", "actes.xlsx"),
    Jeu("groupage", "cma.csv"),
    Jeu("groupage", "racines.xlsx"),
    Jeu("groupage", "tarifs.xlsx", feuille="Tarifs public"),
)

# Liste des CMA livrée par l'ATIH en csv (point-virgule, Windows-1252),
# codes CIM-10 sans point (« C169+0 ») et en-tête abrégé (« diag ; niv ;
# libellé ») : relue sous les noms et la graphie des autres jeux, pour que le
# site la croise code pour code avec les listes de la fonction groupage.
COLONNES_CSV = {"cma.csv": ("Code", "Niveau", "Libellé")}


def code_cim(code: str) -> str:
    """« C169+0 » → « C16.9+0 », « B24+0 » → « B24.+0 », « A09 » → « A09 »."""
    code = code.strip()
    return code if len(code) <= 3 else f"{code[:3]}.{code[3:]}"


def lire_csv(chemin: Path) -> list[dict]:
    colonnes = COLONNES_CSV[chemin.name]
    texte = chemin.read_bytes().decode("cp1252")
    lignes = list(csv.reader(texte.splitlines(), delimiter=";"))[1:]
    resultat = []
    for ligne in lignes:
        if not any(v.strip() for v in ligne):
            continue
        code, niveau, libelle = (v.strip() for v in ligne[:3])
        resultat.append({colonnes[0]: code_cim(code), colonnes[1]: int(niveau), colonnes[2]: libelle})
    return resultat

# Colonnes techniques dont un null xlsx doit se lire comme une chaîne vide et
# non comme une absence de valeur : xlsxwriter ne distingue pas une cellule
# vide d'une cellule nulle, donc le comblement se fait à la lecture.
COLONNES_TEXTE_VIDE_SI_NULLE = {"_caracteristiques"}

# Arrêté tarifaire MCO, tel que le publie l'ATIH : une soixantaine de
# feuilles (GHS, suppléments, forfaits, secteurs public et privé). Trois
# lignes de titre y précèdent l'en-tête, dont les intitulés courent sur
# plusieurs lignes (« TARIF\n(en euros) ») : ils sont relus sous des noms
# courts. Un intitulé ajouté, ôté ou déplacé par l'ATIH arrête la conversion,
# plutôt que de publier une colonne de montants sous le nom d'une autre.
LIGNE_ENTETE = {"tarifs.xlsx": 4}
COLONNES_XLSX = {
    "tarifs.xlsx": {
        "GHS": "GHS",
        "GHM": "GHM",
        "LIBELLE": "Libellé",
        "Bornes basses": "Borne basse",
        "Bornes hautes": "Borne haute",
        "TARIF (en euros)": "Tarif",
        "FORFAIT EXB": "Forfait EXB",
        "TARIF EXB (en euros)": "Tarif EXB",
        "TARIF EXH (en euros)": "Tarif EXH",
    }
}


class ErreurDonnees(Exception):
    pass


def _valeur(v: object) -> object:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def renommer(fichier: str, entetes: list[str]) -> list[str]:
    """Intitulés de l'ATIH → noms courts de COLONNES_XLSX, dans le même ordre."""
    attendus = COLONNES_XLSX[fichier]
    lus = [" ".join(e.split()) for e in entetes]
    while lus and not lus[-1]:  # cellules vides au bout de la ligne d'en-tête
        lus.pop()
    if lus != list(attendus):
        raise ErreurDonnees(f"{fichier} : en-tête inattendu {lus}, attendu {list(attendus)}")
    return [attendus[e] for e in lus]


def lire(chemin: Path, nom_feuille: str | None = None) -> list[dict]:
    """Une ligne par ligne non vide de la feuille (la feuille active par
    défaut), clefs = en-têtes."""
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    try:
        if nom_feuille is None:
            feuille = classeur.active
        elif nom_feuille in classeur.sheetnames:
            feuille = classeur[nom_feuille]
        else:
            raise ErreurDonnees(f"{chemin.name} : pas de feuille « {nom_feuille} » ({', '.join(classeur.sheetnames)})")
        verifie = chemin.name in COLONNES_XLSX
        if verifie:
            # En lecture seule, openpyxl s'arrête à l'étendue que le classeur
            # déclare (balise <dimension>) : fausse, elle tronquerait la
            # feuille sans rien dire. On lit jusqu'à la dernière cellule.
            feuille.reset_dimensions()
        ligne_entete = LIGNE_ENTETE.get(chemin.name, 1)
        lignes = feuille.iter_rows(min_row=ligne_entete, values_only=True)
        entetes = [str(c).strip() if c is not None else "" for c in next(lignes)]
        if verifie:
            entetes = renommer(chemin.name, entetes)
        resultat = []
        for n, ligne in enumerate(lignes, start=ligne_entete + 1):
            if all(v is None for v in ligne):
                continue
            if verifie:
                # Sans étendue déclarée, une ligne s'arrête à sa dernière
                # cellule remplie : on la complète, et une valeur au-delà de
                # l'en-tête arrête la conversion au lieu d'être ignorée.
                if any(v is not None for v in ligne[len(entetes):]):
                    raise ErreurDonnees(f"{chemin.name}, ligne {n} : valeur hors des colonnes de l'en-tête")
                ligne = (*ligne, *[None] * (len(entetes) - len(ligne)))
            enregistrement = {}
            for entete, v in zip(entetes, ligne):
                if entete in COLONNES_TEXTE_VIDE_SI_NULLE and v is None:
                    v = ""
                enregistrement[entete] = _valeur(v)
            resultat.append(enregistrement)
        return resultat
    finally:
        classeur.close()


RE_GHM = re.compile(r"\d{2}[CKMZ]\d{2}[0-9A-Z]")  # en fullmatch : pas de « \n » final
MONTANTS = ("Tarif", "Forfait EXB", "Tarif EXB", "Tarif EXH")

# Racines de la classification qui n'ont aucun GHS dans « Tarifs public ».
# Toute autre racine sans tarif arrête la conversion : un arrêté plus ancien
# que racines.xlsx laisserait sinon ses nouvelles racines sans tarif, en
# silence. Vérifier l'arrêté avant d'ajouter une racine ici.
RACINES_SANS_TARIF = {
    "14Z08",  # Interruptions volontaires de grossesse : séjours de moins de 3 jours
    "15Z10",  # Mort-nés
    "23Z03",  # Interventions de confort et autres interventions non prises en charge par l'AMO
}


def verifier_tarifs(lignes: list[dict]) -> None:
    """Chaque ligne se lit sans rien deviner : un couple GHS-GHM unique, des
    bornes et montants lisibles et cohérents entre eux, un seul tarif par
    GHS (le même, quel que soit le GHM qu'il couvre), un seul libellé par
    GHM. Et l'arrêté couvre la classification publiée : aucune racine
    inconnue de racines.xlsx, aucune racine sans tarif hors de
    RACINES_SANS_TARIF."""
    if not lignes:
        raise ErreurDonnees("tarifs.xlsx : aucune ligne de tarif")
    couples: set[tuple[int, str]] = set()
    par_ghs: dict[int, tuple[str, tuple]] = {}  # GHS → (premier GHM lu, bornes et montants)
    libelles: dict[str, str] = {}
    for ligne in lignes:
        ghs, ghm, libelle = ligne["GHS"], ligne["GHM"], ligne["Libellé"]
        bb, bh = ligne["Borne basse"], ligne["Borne haute"]
        ou = f"tarifs.xlsx, GHS {ghs!r}, GHM {ghm!r}"
        if isinstance(ghs, bool) or not isinstance(ghs, int) or ghs <= 0:
            raise ErreurDonnees(f"{ou} : numéro de GHS illisible")
        if not isinstance(ghm, str) or not RE_GHM.fullmatch(ghm):
            raise ErreurDonnees(f"{ou} : code GHM illisible")
        if not isinstance(libelle, str) or not libelle.strip():
            raise ErreurDonnees(f"{ou} : libellé vide")
        for borne in (bb, bh):
            if isinstance(borne, bool) or not isinstance(borne, int) or borne < 0:
                raise ErreurDonnees(f"{ou} : borne illisible ({borne!r})")
        for colonne in MONTANTS:
            v = ligne[colonne]
            if isinstance(v, bool) or not isinstance(v, (int, float)) or v < 0 or round(v, 2) != v:
                raise ErreurDonnees(f"{ou} : {colonne} illisible ({v!r})")
        if ligne["Tarif"] == 0:
            raise ErreurDonnees(f"{ou} : tarif nul")
        # Une borne à 0 veut dire « pas d'extrême de ce côté » : un montant
        # d'extrême sans sa borne ne s'appliquerait à aucun séjour.
        if bh and bh <= bb:
            raise ErreurDonnees(f"{ou} : borne haute {bh} inférieure ou égale à la borne basse {bb}")
        if ligne["Tarif EXH"] and not bh:
            raise ErreurDonnees(f"{ou} : tarif EXH sans borne haute")
        if (ligne["Tarif EXB"] or ligne["Forfait EXB"]) and not bb:
            raise ErreurDonnees(f"{ou} : extrême bas sans borne basse")
        if ligne["Tarif EXB"] and ligne["Forfait EXB"]:
            raise ErreurDonnees(f"{ou} : forfait EXB et tarif EXB à la fois")
        if (ghs, ghm) in couples:
            raise ErreurDonnees(f"{ou} : couple en double")
        couples.add((ghs, ghm))
        valeurs = (bb, bh, *(ligne[c] for c in MONTANTS))
        autre, attendues = par_ghs.setdefault(ghs, (ghm, valeurs))
        if attendues != valeurs:
            raise ErreurDonnees(f"{ou} : bornes et montants {valeurs} différents de ceux du GHM {autre} {attendues}")
        attendu = libelles.setdefault(ghm, libelle)
        if attendu != libelle:
            raise ErreurDonnees(f"{ou} : libellé {libelle!r} différent de {attendu!r}")
    # Un arrêté d'un autre millésime que la classification publiée : ses
    # racines nouvelles n'auraient ni libellé ni place dans l'arbre.
    racines = {l["ListeRacineGHM"] for l in lire(DOSSIER_DONNEES / "groupage" / "racines.xlsx")}
    tarifees = {ghm[:5] for ghm in libelles}
    inconnues = sorted(tarifees - racines)
    if inconnues:
        raise ErreurDonnees(f"tarifs.xlsx : {len(inconnues)} racine(s) absente(s) de racines.xlsx {inconnues[:10]}")
    sans_tarif = sorted(racines - tarifees - RACINES_SANS_TARIF)
    if sans_tarif:
        raise ErreurDonnees(f"tarifs.xlsx : {len(sans_tarif)} racine(s) de racines.xlsx sans aucun GHS {sans_tarif[:10]}")


# Contrôles propres à un jeu, après lecture et avant écriture du JSON.
VERIFICATIONS = {"tarifs.xlsx": verifier_tarifs}

# Campagne de l'arrêté tarifaire, que le classeur de l'ATIH ne nomme nulle
# part (ni titre, ni propriétés, ni feuille) : elle est déclarée ici, avec
# l'empreinte SHA-256 du classeur auquel elle se rapporte. Remplacer
# tarifs.xlsx sans mettre cette ligne à jour arrête la conversion, plutôt
# que de publier les tarifs d'une campagne sous l'année d'une autre.
CAMPAGNES = {
    "tarifs.xlsx": (2026, "f47995e2fb8ae9a846ca6119a10317f4c1fe2ff4152742a34185ef28521dda09"),
}


def campagne(source: Path) -> int:
    annee, empreinte = CAMPAGNES[source.name]
    lue = hashlib.sha256(source.read_bytes()).hexdigest()
    if lue != empreinte:
        raise ErreurDonnees(
            f"{source.name} a changé : déclarer sa campagne dans CAMPAGNES, avec son empreinte {lue}"
        )
    return annee


def convertir(jeu: Jeu) -> None:
    source = DOSSIER_DONNEES / jeu.theme / jeu.fichier
    lignes = lire_csv(source) if source.suffix == ".csv" else lire(source, jeu.feuille)
    if jeu.fichier in VERIFICATIONS:
        VERIFICATIONS[jeu.fichier](lignes)

    dossier_sortie = DOSSIER_SORTIE / jeu.theme
    dossier_sortie.mkdir(parents=True, exist_ok=True)
    cible = dossier_sortie / f"{source.stem}.json"

    # Format colonnaire (en-têtes une fois + lignes en tableaux) plutôt
    # qu'un objet par ligne : ces jeux répètent jusqu'à 19 clefs identiques
    # sur des milliers de lignes, ce qui gonflait le JSON transféré au
    # navigateur de plus de moitié pour rien. `donnees.js` reconstitue les
    # objets à la volée au chargement, donc rien d'autre ne change côté site.
    colonnes = list(lignes[0].keys()) if lignes else []
    valeurs = [[ligne[c] for c in colonnes] for ligne in lignes]
    payload = {"millesime": millesime(source)}
    if source.name in CAMPAGNES:
        payload["campagne"] = campagne(source)
    payload |= {"colonnes": colonnes, "valeurs": valeurs}
    cible.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(f"{cible.relative_to(RACINE)} : {len(lignes)} lignes")


def main() -> None:
    try:
        for jeu in JEUX:
            convertir(jeu)
    except ErreurDonnees as e:
        sys.exit(f"Conversion impossible : {e}")


if __name__ == "__main__":
    main()
